"""Excel copy of the rulebook (generated from rulebook/v1.yaml so the spreadsheet always matches the code)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER = PatternFill('solid', fgColor='0F3D32')


def describe(condition, depth=0):
    """A readable form of one condition: `a >= b`, or `any of:` / `all of:` with indented lines."""
    if isinstance(condition, dict):
        (kind, items), = condition.items()
        inner = '\n'.join('  '*(depth+1)+'- '+describe(i, depth+1) for i in items)
        return f'{kind} of:\n{inner}'
    left, op, right = condition
    return f'{left} {op} {right}'


def _sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color='FFFFFF'), HEADER
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'B2'
    return ws


def write(path, book):
    wb = Workbook()
    about = wb.active
    about.title = 'About'
    for row in (('Name', book.get('name')), ('Version', book['version']), ('Hash', book['hash']),
                ('Source', 'rulebook/v1.yaml (generated; edit the YAML, not this file)'),
                ('Paper only', 'Nothing here places or sizes real trades.'),
                *[(f'Default: {k}', v) for k, v in book['defaults'].items()],
                *[(f'Layer: {k}', v) for k, v in book['layers'].items()]):
        about.append(list(row))
    about.column_dimensions['A'].width, about.column_dimensions['B'].width = 32, 70
    lines = lambda items: '\n'.join(describe(c) for c in items or [])
    _sheet(wb, 'Playbooks', ['ID', 'Name', 'Family', 'Side', 'Owner rules', 'Thesis', 'Trigger', 'Regimes', 'Conditions',
                             'Stop', 'Target', 'Abort', 'Time stop (h)', 'Break-even at R', 'May flip', 'Priority',
                             'Shadow only'],
           [[p['id'], p['name'], p.get('family'), p['side'], ', '.join(p.get('source', [])), p.get('thesis'),
             f"{' / '.join(p['trigger']['any'])} within {p['trigger']['within_bars']} bars "
             f"(after at least {p['trigger']['min_bars']})", ', '.join(p['regimes']), lines(p.get('conditions')),
             f"{p['stop']['anchor']} {'-' if p['side'] == 'long' else '+'} {p['stop'].get('buffer_atr', 0)} ATR",
             f"nearest of {', '.join(p['target'].get('candidates', [])) or 'none'}; cap {p['target'].get('cap')}; "
             f"fallback {p['target'].get('fallback_r')}R", lines(p.get('abort')), p['time_stop_hours'],
             p.get('breakeven_at_r'), bool(p.get('may_flip')), p['priority'], bool(p.get('shadow_only'))]
            for p in book['playbooks']],
           [8, 30, 12, 7, 18, 50, 34, 30, 48, 26, 40, 34, 10, 10, 8, 8, 8])
    _sheet(wb, 'Triggers', ['Trigger', 'Direction', 'Evaluated', 'Condition'],
           [[name, spec['direction'], spec.get('at', spec.get('source')), describe(spec['when']) if 'when' in spec
             else f"from the {spec.get('source')} records"] for name, spec in book['triggers'].items()],
           [26, 10, 10, 80])
    _sheet(wb, 'Gates', ['Gate', 'Setting'], [[k, str(v)] for k, v in book['gates'].items()], [26, 80])
    _sheet(wb, 'Confluence', ['Factor', 'Vote'], [[c['factor'], c['rule']] for c in book['confluence']], [22, 80])
    _sheet(wb, 'Diagnostics', ['Code', 'Meaning'], [[k, v] for k, v in book['diagnostics'].items()], [24, 80])
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
