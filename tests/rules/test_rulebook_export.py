"""The Excel copy of the rulebook is generated from rulebook/v1.yaml, never edited by hand."""
from openpyxl import load_workbook

from quantpilot.rules import export as rulebook_export
from quantpilot.rules import rulebook as rules


def test_workbook_mirrors_the_rulebook(tmp_path):
    book = rules.load()
    path = rulebook_export.write(tmp_path/'rulebook.xlsx', book)
    wb = load_workbook(path)
    assert wb.sheetnames == ['About', 'Playbooks', 'Triggers', 'Gates', 'Confluence', 'Diagnostics']
    playbooks = wb['Playbooks']
    header = [c.value for c in playbooks[1]]
    assert header[:4] == ['ID', 'Name', 'Family', 'Side'] and 'Conditions' in header
    assert playbooks.max_row == len(book['playbooks'])+1
    ids = [playbooks.cell(row=r, column=1).value for r in range(2, playbooks.max_row+1)]
    assert ids == [p['id'] for p in book['playbooks']]
    conditions = playbooks.cell(row=2, column=header.index('Conditions')+1).value
    assert 'trigger.oi_change <= -0.005' in conditions and 'any of' in conditions
    about = {r[0].value: r[1].value for r in wb['About'].iter_rows()}
    assert about['Hash'] == book['hash'] and about['Version'] == book['version']
    assert wb['Triggers'].max_row == len(book['triggers'])+1
